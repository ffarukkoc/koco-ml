import json
from contextlib import contextmanager
from datetime import date

import joblib
import pandas as pd
import pytest
from openpyxl import load_workbook

import koco
from _samples import BinaryTable, RegressionTable, WOETable

Stamp = date.today().strftime("%y%m%d")

# ---------- JSON ----------

def test_json_round_trip_and_default_name(tmp_path, monkeypatch):
    Plan = {"Columns": {"a": {"Strategy": "Median", "Value": 1.5}}}
    assert koco.ExportJson(Plan, tmp_path / "sub" / "plan.json") == tmp_path / "sub" / "plan.json"
    assert koco.ImportJson(tmp_path / "sub" / "plan.json") == Plan
    monkeypatch.chdir(tmp_path)
    assert koco.ExportJson(Plan) == __import__("pathlib").Path(f"{Stamp}_ExportJson.json")
    assert koco.ImportJson(f"{Stamp}_ExportJson.json") == Plan
    assert koco.ExportJson(Plan, "Plans/Nested") == __import__("pathlib").Path("Plans/Nested") / f"{Stamp}_ExportJson.json"

def test_export_reports_an_existing_file(tmp_path, capsys):
    koco.ExportJson({"a": 1}, tmp_path / "x.json")
    capsys.readouterr()
    koco.ExportJson({"a": 2}, tmp_path / "x.json")
    assert "same name" in capsys.readouterr().out
    assert koco.ImportJson(tmp_path / "x.json") == {"a": 2}

def test_export_json_rejects_nan_and_non_dict(tmp_path):
    with pytest.raises(ValueError):
        koco.ExportJson({"Value": float("nan")}, tmp_path / "bad.json")
    assert not (tmp_path / "bad.json").exists()
    with pytest.raises(TypeError):
        koco.ExportJson([1, 2], tmp_path / "list.json")
    with pytest.raises(TypeError, match="FilePath"):
        koco.ExportJson({"a": 1}, 5)

def test_import_json_requires_an_object(tmp_path):
    (tmp_path / "list.json").write_text("[1, 2]", encoding="utf-8")
    with pytest.raises(ValueError):
        koco.ImportJson(tmp_path / "list.json")

def test_import_excel_round_trip(tmp_path):
    First = pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", None], "c": [0.5, 1.5, 2.5]})
    Second = pd.DataFrame({"d": [10, 20]})
    Path = koco.ExportExcel({"One": First, "Two": Second}, tmp_path / "book.xlsx")
    pd.testing.assert_frame_equal(koco.ImportExcel(Path), First)
    pd.testing.assert_frame_equal(koco.ImportExcel(Path, "Two"), Second)
    pd.testing.assert_frame_equal(koco.ImportExcel(Path, SheetName=1), Second)
    Everything = koco.ImportExcel(Path, SheetName=None)
    assert list(Everything) == ["One", "Two"]
    pd.testing.assert_frame_equal(Everything["One"], First)

def test_import_excel_errors(tmp_path):
    with pytest.raises(FileNotFoundError, match="not found"):
        koco.ImportExcel(tmp_path / "missing.xlsx")
    Path = koco.ExportExcel({"One": pd.DataFrame({"a": [1]})}, tmp_path / "book.xlsx")
    with pytest.raises(ValueError):
        koco.ImportExcel(Path, "NoSuchSheet")

# ---------- Parquet / Joblib ----------

def test_parquet_round_trip(tmp_path, monkeypatch):
    Table = pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})
    Path = koco.ExportParquet(Table, tmp_path / "t" / "table.parquet")
    pd.testing.assert_frame_equal(koco.ImportParquet(Path), Table)
    monkeypatch.chdir(tmp_path)
    assert koco.ExportParquet(Table).name == f"{Stamp}_ExportParquet.parquet"

def test_joblib_round_trip(tmp_path, monkeypatch):
    Path = koco.ExportJoblib({"a": [1, 2]}, tmp_path / "x" / "d.joblib")
    assert koco.ImportJoblib(Path) == {"a": [1, 2]}
    monkeypatch.chdir(tmp_path)
    assert koco.ExportJoblib({"a": 1}).name == f"{Stamp}_ExportJoblib.joblib"

# ---------- Excel ----------

def test_export_to_excel_writes_every_table(tmp_path):
    Path = koco.ExportExcel(
        {"One": pd.DataFrame({"a": [1, 2]}), "Two": [pd.DataFrame({"b": [3]}), pd.DataFrame({"b": [4]})]},
        tmp_path / "report.xlsx",
    )
    Workbook = load_workbook(Path)
    assert Workbook.sheetnames == ["One", "Two"]
    assert Workbook["One"]["A2"].value == 1
    assert Workbook["Two"]["A5"].value == 4

def test_export_to_excel_default_and_folder_names(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    Sheets = {"One": pd.DataFrame({"a": [1]})}
    assert koco.ExportExcel(Sheets).name == f"{Stamp}_ExportExcel.xlsx"
    assert koco.ExportExcel(Sheets, "Reports").parent.name == "Reports"
    with pytest.raises(TypeError, match="argument order"):
        koco.ExportExcel("report.xlsx", Sheets)

def test_export_to_excel_validates_before_writing(tmp_path, monkeypatch):
    Path = tmp_path / "report.xlsx"
    Table = pd.DataFrame({"a": [1]})
    for Sheets, Error in [
        ({}, ValueError),
        ({"bad/name": Table}, ValueError),
        ({"x" * 32: Table}, ValueError),
        ({"Same": Table, "same": Table}, ValueError),
        ({"Ok": 5}, TypeError),
    ]:
        with pytest.raises(Error):
            koco.ExportExcel(Sheets, Path)
    monkeypatch.setattr("koco.general._EXCEL_MAX_ROWS", 3)
    with pytest.raises(ValueError, match="limit"):
        koco.ExportExcel({"Big": pd.DataFrame({"a": range(5)})}, Path)
    assert not Path.exists()

# ---------- Misc ----------

def test_run_func_in_subsample_adds_group_column():
    Table = pd.DataFrame({"g": ["x", "x", "y"], "v": [1, 2, 3]})
    Result = koco.RunFuncInSubSample(Table, "g", lambda Sub: pd.DataFrame({"n": [len(Sub)]}))
    assert Result.to_dict("list") == {"g": ["x", "y"], "n": [2, 1]}
    with pytest.raises(ValueError, match="not found"):
        koco.RunFuncInSubSample(Table, "zz", lambda Sub: Sub)
    with pytest.raises(TypeError):
        koco.RunFuncInSubSample(Table, "g", lambda Sub: 1)

def test_export_model_artifacts_writes_manifest(tmp_path, monkeypatch):
    Table = RegressionTable()
    Artifacts, _, _, _ = koco.RegressionModelling(Table, ["RidgeR"], "y", Split=False)
    Directory = koco.ExportModelArtifacts(Artifacts, "Test models", tmp_path / "models")
    Manifest = json.loads((Directory / "Manifest.json").read_text(encoding="utf-8"))
    assert Manifest["Explanation"] == "Test models"
    assert Manifest["Models"]["RidgeR"]["File"] == "RidgeR.joblib"
    assert Manifest["Models"]["RidgeR"]["ModelType"] == "Regression"
    assert len(koco.ApplyModel(Table, koco.ImportJoblib(Directory / "RidgeR.joblib"))) == len(Table)
    monkeypatch.chdir(tmp_path / "models")
    koco.ExportModelArtifacts(Artifacts)
    assert (tmp_path / "models" / "RidgeR.joblib").exists()
    with pytest.raises(TypeError):
        koco.ExportModelArtifacts(Artifacts, DirPath=5)

# ---------- Connection ----------

class _FakeEngine:
    @contextmanager
    def connect(self):
        yield self

def _FakeCreateEngine(Captured):
    def Create(Url, **Options):
        Captured.append(Url)
        return _FakeEngine()
    return Create

SqlConfig = {
    "DefaultServer": "SERVER_A",
    "DefaultDatabase": "MyDatabase",
    "DefaultDriver": "My Driver",
    "Trusted_Connection": "yes",
    "SPNs": {"SERVER_A": "MSSQLSvc/server-a.example.com:instance", "SERVER_B": "MSSQLSvc/server-b.example.com"},
}

def _WriteConfig(tmp_path, Config=SqlConfig, Name="sql.json"):
    ConfigPath = tmp_path / Name
    ConfigPath.write_text(json.dumps(Config), encoding="utf-8")
    return ConfigPath

def _CaptureUrls(monkeypatch, System="Windows"):
    Urls = []
    monkeypatch.setattr("sqlalchemy.create_engine", _FakeCreateEngine(Urls))
    monkeypatch.setattr("koco.general._SetKerberosCache", lambda: None)
    monkeypatch.setattr("platform.system", lambda: System)
    return Urls

def test_connection_needs_server_and_database_without_a_config(monkeypatch):
    Urls = _CaptureUrls(monkeypatch)
    for Arguments in ({}, {"Server": "SERVER_A"}, {"Database": "MyDatabase"}):
        with pytest.raises(ValueError, match="must be given") as Info:
            koco.Connection(**Arguments)
        assert "DefaultServer" in str(Info.value)
    assert Urls == []

def test_connection_without_a_config_uses_the_built_in_defaults(monkeypatch):
    Urls = _CaptureUrls(monkeypatch)
    koco.Connection("SERVER_A", "MyDatabase")
    assert Urls[0].startswith("mssql+pyodbc:///?odbc_connect=")
    assert "ODBC+Driver+17+for+SQL+Server" in Urls[0] and "Trusted_Connection%3Dyes" in Urls[0]
    assert "ServerSPN" not in Urls[0]

def test_connection_reads_defaults_from_the_config(tmp_path, monkeypatch):
    Urls = _CaptureUrls(monkeypatch)
    ConfigPath = _WriteConfig(tmp_path)
    koco.Connection(Config=ConfigPath)
    koco.Connection("SERVER_B", Config=str(ConfigPath))
    koco.Connection(Database="OtherDb", Config=ConfigPath)
    koco.Connection("SERVER_B", "OtherDb", ConfigPath)
    assert "SERVER%3DSERVER_A" in Urls[0] and "DATABASE%3DMyDatabase" in Urls[0] and "My+Driver" in Urls[0]
    assert "SERVER%3DSERVER_B" in Urls[1] and "DATABASE%3DMyDatabase" in Urls[1]
    assert "SERVER%3DSERVER_A" in Urls[2] and "DATABASE%3DOtherDb" in Urls[2]
    assert "SERVER%3DSERVER_B" in Urls[3] and "DATABASE%3DOtherDb" in Urls[3]

def test_connection_arguments_are_never_read_from_the_config(tmp_path, monkeypatch):
    Urls = _CaptureUrls(monkeypatch)
    ConfigPath = _WriteConfig(tmp_path, {"DefaultDatabase": "FromFile"})
    koco.Connection("SERVER_A", "FromCall", ConfigPath)
    assert "FromCall" in Urls[0] and "FromFile" not in Urls[0]

def test_connection_config_without_defaults_still_needs_names(tmp_path, monkeypatch):
    _CaptureUrls(monkeypatch)
    ConfigPath = _WriteConfig(tmp_path, {"SPNs": {"SERVER_A": "x"}})
    with pytest.raises(ValueError, match="DefaultServer and DefaultDatabase"):
        koco.Connection(Config=ConfigPath)
    with pytest.raises(ValueError, match="DefaultDatabase"):
        koco.Connection("SERVER_A", Config=ConfigPath)

def test_connection_config_driver_and_trusted_connection(tmp_path, monkeypatch):
    Urls = _CaptureUrls(monkeypatch)
    koco.Connection("S", "D", _WriteConfig(tmp_path, {"DefaultDriver": "Other Driver", "Trusted_Connection": "NO"}))
    koco.Connection("S", "D", _WriteConfig(tmp_path, {}, "empty.json"))
    assert "Other+Driver" in Urls[0] and "Trusted_Connection%3Dno" in Urls[0]
    assert "ODBC+Driver+17" in Urls[1] and "Trusted_Connection%3Dyes" in Urls[1]

def test_connection_missing_config_file_gives_a_clear_error(tmp_path):
    with pytest.raises(FileNotFoundError, match="configuration file not found") as Info:
        koco.Connection("S", "D", tmp_path / "missing.json")
    assert '"SPNs"' in str(Info.value) and "without Config" in str(Info.value)

def test_connection_expands_the_home_folder(tmp_path, monkeypatch):
    Urls = _CaptureUrls(monkeypatch)
    _WriteConfig(tmp_path, Name="home_sql.json")
    monkeypatch.setenv("HOME", str(tmp_path))
    koco.Connection(Config="~/home_sql.json")
    assert "SERVER_A" in Urls[0]

def test_connection_rejects_an_invalid_config(tmp_path, monkeypatch):
    _CaptureUrls(monkeypatch)
    (tmp_path / "bad.json").write_text("{not json", encoding="utf-8")
    with pytest.raises(ValueError, match="not valid JSON"):
        koco.Connection("S", "D", tmp_path / "bad.json")
    for Index, Config in enumerate((
        ["x"], {"DefaultServer": 5}, {"DefaultDriver": []}, {"Trusted_Connection": "maybe"}, {"Trusted_Connection": True},
        {"SPNs": ["x"]}, {"SPNs": {"S": 5}}, {"ODBCSYSINI": 5}, {"KRB5_CONFIG": ""},
    )):
        with pytest.raises(ValueError):
            koco.Connection("S", "D", _WriteConfig(tmp_path, Config, f"invalid{Index}.json"))
    for Wrong in ({"DefaultServer": "S"}, 5, ["a.json"]):
        with pytest.raises(TypeError, match="Config"):
            koco.Connection("S", "D", Wrong)

def test_connection_on_macos_uses_the_spn_of_the_server(tmp_path, monkeypatch):
    Urls = _CaptureUrls(monkeypatch, System="Darwin")
    ConfigPath = _WriteConfig(tmp_path)
    koco.Connection("SERVER_B", "MyDatabase", ConfigPath)
    assert "ServerSPN%3DMSSQLSvc%2Fserver-b.example.com" in Urls[0]
    assert "MARS_Connection%3Dyes" in Urls[0] and "TrustServerCertificate%3Dyes" in Urls[0]

def test_connection_on_macos_needs_an_spn(tmp_path, monkeypatch):
    Urls = _CaptureUrls(monkeypatch, System="Darwin")
    with pytest.raises(ValueError, match="needs an SPN") as Info:
        koco.Connection("SERVER_A", "MyDatabase")
    assert "pass Config" in str(Info.value)
    with pytest.raises(ValueError, match="needs an SPN"):
        koco.Connection("UnknownServer", "MyDatabase", _WriteConfig(tmp_path))
    assert Urls == []

def _ClearSqlEnvironment(monkeypatch):
    # setenv first so that monkeypatch removes the variables the code sets.
    for Key in ("ODBCSYSINI", "KRB5_CONFIG"):
        monkeypatch.setenv(Key, "placeholder")
        monkeypatch.delenv(Key)

def _WriteSqlEnvironmentFiles(tmp_path):
    (tmp_path / "odbc").mkdir()
    (tmp_path / "odbc" / "odbcinst.ini").write_text("[ODBC Driver 17 for SQL Server]\n", encoding="utf-8")
    (tmp_path / "krb5.conf").write_text("[libdefaults]\n", encoding="utf-8")
    return {**SqlConfig, "ODBCSYSINI": str(tmp_path / "odbc"), "KRB5_CONFIG": "krb5.conf"}

def test_connection_sets_the_odbc_and_kerberos_environment(tmp_path, monkeypatch):
    import os
    _ClearSqlEnvironment(monkeypatch)
    Urls = _CaptureUrls(monkeypatch, System="Darwin")
    koco.Connection("SERVER_A", "MyDatabase", _WriteConfig(tmp_path, _WriteSqlEnvironmentFiles(tmp_path)))
    assert os.environ["ODBCSYSINI"] == str(tmp_path / "odbc")
    assert os.environ["KRB5_CONFIG"] == str((tmp_path / "krb5.conf").resolve())
    assert len(Urls) == 1

def test_connection_keeps_an_existing_environment_variable(tmp_path, monkeypatch, capsys):
    import os
    _ClearSqlEnvironment(monkeypatch)
    monkeypatch.setenv("ODBCSYSINI", "/somewhere/else")
    _CaptureUrls(monkeypatch, System="Linux")
    koco.Connection("SERVER_A", "MyDatabase", _WriteConfig(tmp_path, _WriteSqlEnvironmentFiles(tmp_path)))
    assert os.environ["ODBCSYSINI"] == "/somewhere/else"
    assert "ODBCSYSINI is already set" in capsys.readouterr().out
    assert "KRB5_CONFIG" in os.environ

def test_connection_checks_the_environment_paths(tmp_path, monkeypatch):
    _ClearSqlEnvironment(monkeypatch)
    Urls = _CaptureUrls(monkeypatch, System="Darwin")
    (tmp_path / "empty").mkdir()
    with pytest.raises(FileNotFoundError, match="odbcinst.ini"):
        koco.Connection("SERVER_A", "D", _WriteConfig(tmp_path, {**SqlConfig, "ODBCSYSINI": str(tmp_path / "empty")}))
    with pytest.raises(FileNotFoundError, match="KRB5_CONFIG"):
        koco.Connection("SERVER_A", "D", _WriteConfig(tmp_path, {**SqlConfig, "KRB5_CONFIG": "missing.conf"}, "k.json"))
    assert Urls == []

def test_connection_ignores_the_environment_keys_on_windows(tmp_path, monkeypatch):
    import os
    _ClearSqlEnvironment(monkeypatch)
    Urls = _CaptureUrls(monkeypatch)
    koco.Connection("S", "D", _WriteConfig(tmp_path, {"ODBCSYSINI": "/missing", "KRB5_CONFIG": "/missing.conf"}))
    assert "ODBCSYSINI" not in os.environ and "KRB5_CONFIG" not in os.environ
    assert len(Urls) == 1

# ---------- Model bundle ----------

def _FittedBundleParts():
    Table = RegressionTable()
    Table.loc[::10, "a"] = None
    NullPlan, _ = koco.FitNullPlan(Table, [("a", "Median")])
    Filled = koco.ApplyNullPlan(Table, NullPlan, Verbose=False)
    OutlierPlan = {"Columns": {"b": ["U", 1.0]}}
    Capped = koco.ApplyOutlierPlan(Filled, OutlierPlan)
    ScalePlan = koco.FitScalingPlan(Capped, {"Standardization": ["a", "b"]})
    Scaled = koco.ApplyScalingPlan(Capped, ScalePlan)
    Artifacts, _, _, _ = koco.RegressionModelling(Scaled, ["RidgeR"], "y", Split=False)
    return Table, Scaled, Artifacts["RidgeR"], {"Null": NullPlan, "Outlier": OutlierPlan, "Scaling": ScalePlan}

def test_model_bundle_scores_raw_data_like_the_manual_pipeline(tmp_path):
    Raw, Scaled, Artifact, Plans = _FittedBundleParts()
    Path, Digest = koco.ExportModelBundle(Artifact, Plans, tmp_path / "bundle" / "model.joblib", "Test bundle")
    Bundle = koco.ImportModelBundle(Path, ExpectedHash=Digest)
    assert set(Bundle["Plans"]) == {"Null", "Outlier", "Scaling"} and Bundle["Explanation"] == "Test bundle"
    Expected = koco.ApplyModel(Scaled, Artifact)
    pd.testing.assert_series_equal(koco.ApplyModelBundle(Raw, Bundle), Expected)

def test_model_bundle_plans_are_optional(tmp_path):
    Table = BinaryTable()
    Artifacts, _, _, _ = koco.ClassificationModelling(Table, ["LogisticC"], "y", Split=False)
    Path, _ = koco.ExportModelBundle(Artifacts["LogisticC"], FilePath=tmp_path)
    assert Path.name == f"{Stamp}_ExportModelBundle.joblib"
    Bundle = koco.ImportModelBundle(Path)
    assert Bundle["Plans"] == {}
    Scores = koco.ApplyModelBundle(Table, Bundle)
    pd.testing.assert_frame_equal(Scores, koco.ApplyModel(Table, Artifacts["LogisticC"]))
    OnlyScaling, _ = koco.ExportModelBundle(Artifacts["LogisticC"], {"Scaling": koco.FitScalingPlan(Table, {"Robust": ["a"]})}, tmp_path / "s.joblib")
    assert list(koco.ImportModelBundle(OnlyScaling)["Plans"]) == ["Scaling"]

def test_model_bundle_integrity_checks(tmp_path):
    _, _, Artifact, Plans = _FittedBundleParts()
    Path, Digest = koco.ExportModelBundle(Artifact, Plans, tmp_path / "m.joblib")
    with pytest.raises(ValueError, match="ExpectedHash"):
        koco.ImportModelBundle(Path, ExpectedHash="0" * 64)
    Stored = joblib.load(Path)
    Stored["ModelSha256"] = "0" * 64
    joblib.dump(Stored, tmp_path / "tampered.joblib")
    with pytest.raises(ValueError, match="corrupted"):
        koco.ImportModelBundle(tmp_path / "tampered.joblib")
    joblib.dump({"a": 1}, tmp_path / "other.joblib")
    with pytest.raises(ValueError, match="not a koco model bundle"):
        koco.ImportModelBundle(tmp_path / "other.joblib")

def test_model_bundle_export_validation(tmp_path):
    _, _, Artifact, Plans = _FittedBundleParts()
    with pytest.raises(ValueError, match="Unknown plan names"):
        koco.ExportModelBundle(Artifact, {"Imputation": Plans["Null"]}, tmp_path / "a.joblib")
    with pytest.raises(TypeError):
        koco.ExportModelBundle(Artifact, [Plans["Null"]], tmp_path / "a.joblib")
    with pytest.raises(ValueError, match="Columns"):
        koco.ExportModelBundle(Artifact, {"Null": {"x": 1}}, tmp_path / "a.joblib")
    with pytest.raises(ValueError, match="ModelType"):
        koco.ExportModelBundle({"Model": 1, "Features": ["a"]}, None, tmp_path / "a.joblib")
    with pytest.raises(ValueError, match="ImportModelBundle"):
        koco.ApplyModelBundle(RegressionTable(), {"a": 1})
    assert not (tmp_path / "a.joblib").exists()

def test_model_bundle_with_a_woe_plan_scores_raw_data(tmp_path):
    Raw = WOETable(Rows=2000)
    NullPlan, _ = koco.FitNullPlan(Raw, [("income", "Median")])
    Filled = koco.ApplyNullPlan(Raw, NullPlan, Verbose=False)
    WOEPlan, _, _ = koco.FitWOEPlan(Filled, "y", Features=["score", "age", "sparse"], CoverBound=0.05, MeanDiffBound=0.02)
    Transformed = koco.ApplyWOEPlan(Filled, WOEPlan)
    Features = [f"{Column}_WOE" for Column in WOEPlan["Columns"]]
    Artifacts, _, _, _ = koco.ClassificationModelling(Transformed, ["LogisticC"], "y", Features=Features, Split=False)

    Path, Digest = koco.ExportModelBundle(Artifacts["LogisticC"], {"WOE": WOEPlan, "Null": NullPlan}, tmp_path / "woe.joblib")
    Bundle = koco.ImportModelBundle(Path, ExpectedHash=Digest)
    assert list(Bundle["Plans"]) == ["Null", "WOE"]
    pd.testing.assert_frame_equal(koco.ApplyModelBundle(Raw, Bundle), koco.ApplyModel(Transformed, Artifacts["LogisticC"]))
