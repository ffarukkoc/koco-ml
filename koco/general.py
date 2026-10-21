from __future__ import annotations

import hashlib
import importlib
import io
import json
import sys
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ._helpers import (
    BUNDLE_PLANS, _Alert, _CheckArtifact, _CheckColumns, _Done, _Info, _OutputPath, _Step, _ValidatePlan,
)

if TYPE_CHECKING:
    from sqlalchemy.engine import Engine

__all__ = [
    "Connection", "RunFuncInSubSample", "ExportExcel", "ImportExcel", "ExportParquet", "ImportParquet",
    "ExportJoblib", "ImportJoblib", "ExportJson", "ImportJson", "ExportModelArtifacts",
    "ExportModelBundle", "ImportModelBundle",
]

#********* General Functions *********#

_EXCEL_MAX_ROWS = 1_048_576
_EXCEL_MAX_COLUMNS = 16_384
_EXCEL_INVALID_SHEET_CHARS = set("[]:*?/\\")

_SQL_DEFAULT_DRIVER = "ODBC Driver 17 for SQL Server"
_SQL_DEFAULT_TRUSTED_CONNECTION = "yes"
_SQL_CONFIG_EXAMPLE = """{
  "DefaultServer": "SERVER_A",
  "DefaultDatabase": "MyDatabase",
  "DefaultDriver": "ODBC Driver 17 for SQL Server",
  "Trusted_Connection": "yes",
  "ODBCSYSINI": "/opt/homebrew/etc",
  "KRB5_CONFIG": "~/krb5.conf",
  "SPNs": {
    "SERVER_A": "MSSQLSvc/server-a.example.com:instance",
    "SERVER_B": "MSSQLSvc/server-b.example.com"
  }
}"""
_SQL_ENVIRONMENT_KEYS = ("ODBCSYSINI", "KRB5_CONFIG")

def _LoadSqlConfig(
    Config: str | Path,
) -> dict:
    """
    Read and validate a SQL Server configuration file.

    Arg:
       Config: Path of the JSON file (see Connection for the keys).

    Returns:
       Dict with DefaultServer, DefaultDatabase, Driver, TrustedConnection, SPNs and Environment;
       missing keys get the built-in defaults (None for the server and database, {} for SPNs and
       Environment). Environment maps ODBCSYSINI / KRB5_CONFIG to paths; a relative path is taken
       from the folder of the file.

    Raises:
       FileNotFoundError: If the file does not exist.
       ValueError: If the file is not valid JSON or a key has the wrong type.
    """
    ConfigPath = Path(Config).expanduser()
    if not ConfigPath.is_file():
        raise FileNotFoundError(
            f"SQL configuration file not found: {ConfigPath}\n"
            f"Create it, or call Connection with Server and Database and without Config. Example content:\n{_SQL_CONFIG_EXAMPLE}\n"
            "Every key is optional. SPNs (server -> service principal name) are needed on macOS only; "
            "ODBCSYSINI and KRB5_CONFIG are used on macOS / Linux only."
        )
    try:
        Raw = json.loads(ConfigPath.read_text(encoding="utf-8"))
    except json.JSONDecodeError as Error:
        raise ValueError(f"The SQL configuration {ConfigPath} is not valid JSON: {Error}") from Error
    if not isinstance(Raw, dict):
        raise ValueError(f"The SQL configuration {ConfigPath} must contain a JSON object.")

    for Key in ("DefaultServer", "DefaultDatabase", "DefaultDriver"):
        if Raw.get(Key) is not None and not isinstance(Raw[Key], str):
            raise ValueError(f'"{Key}" in {ConfigPath} must be a string.')
    Trusted = Raw.get("Trusted_Connection", _SQL_DEFAULT_TRUSTED_CONNECTION)
    if not isinstance(Trusted, str) or Trusted.lower() not in ("yes", "no"):
        raise ValueError(f'"Trusted_Connection" in {ConfigPath} must be "yes" or "no".')
    SPNs = Raw.get("SPNs", {})
    if not isinstance(SPNs, dict) or not all(isinstance(Value, str) for Value in SPNs.values()):
        raise ValueError(f'"SPNs" in {ConfigPath} must map server names to SPN strings.')
    Environment = {}
    for Key in _SQL_ENVIRONMENT_KEYS:
        Value = Raw.get(Key)
        if Value is None:
            continue
        if not isinstance(Value, str) or not Value.strip():
            raise ValueError(f'"{Key}" in {ConfigPath} must be a non-empty path string.')
        EnvPath = Path(Value).expanduser()
        Environment[Key] = EnvPath if EnvPath.is_absolute() else (ConfigPath.parent / EnvPath).resolve()

    return {
        "DefaultServer": Raw.get("DefaultServer"),
        "DefaultDatabase": Raw.get("DefaultDatabase"),
        "Driver": Raw.get("DefaultDriver") or _SQL_DEFAULT_DRIVER,
        "TrustedConnection": Trusted.lower(),
        "SPNs": SPNs,
        "Environment": Environment,
    }

def _SetSqlEnvironment(
    Environment: dict,
) -> None:
    """
    Tell unixODBC and Kerberos where their configuration files are.

    Arg:
       Environment: Dict from _LoadSqlConfig mapping ODBCSYSINI (folder with odbcinst.ini) and
          KRB5_CONFIG (krb5.conf file) to paths.

    Returns:
       None. Sets the environment variables of the current process. A variable that is already set
       is kept; a different value is reported.

    Raises:
       FileNotFoundError: If the ODBCSYSINI folder has no odbcinst.ini or the KRB5_CONFIG file does not exist.
    """
    import os

    for Key, EnvPath in Environment.items():
        Target = EnvPath / "odbcinst.ini" if Key == "ODBCSYSINI" else EnvPath
        if not Target.is_file():
            raise FileNotFoundError(f'"{Key}" of the SQL configuration points to {EnvPath}, but {Target} does not exist.')
        Current = os.environ.get(Key)
        if Current is None:
            os.environ[Key] = str(EnvPath)
        elif Path(Current).expanduser().resolve() != EnvPath.resolve():
            _Info(f"{Key} is already set to {Current}; the value {EnvPath} of the SQL configuration is not used.")

def _SetKerberosCache() -> None:
    """
    Point unixODBC to the macOS API Kerberos credential cache found by klist.

    Arg:
       None.

    Returns:
       None. Sets the KRB5CCNAME environment variable of the current process.

    Raises:
       RuntimeError: If klist is not available or no API credential cache is found.
    """
    import os
    import re
    import subprocess

    try:
        KlistOutput = subprocess.check_output(["klist"], text=True)
    except (OSError, subprocess.CalledProcessError) as Error:
        raise RuntimeError("klist failed; get a Kerberos ticket first (kinit).") from Error
    CacheMatch = re.search(r"(API:[A-F0-9\-]+)", KlistOutput)
    if CacheMatch is None:
        raise RuntimeError("No Kerberos API credential cache found in the klist output; get a ticket first (kinit).")
    os.environ["KRB5CCNAME"] = CacheMatch.group(1)

def Connection(
    Server: str | None = None,
    Database: str | None = None,
    Config: str | Path | None = None,
) -> Engine:
    """
    Create a SQLAlchemy engine for a SQL Server database.

    Server and Database are given directly or come from the DefaultServer and DefaultDatabase of a
    JSON configuration file. A value given to the function is never read from the file. The file
    can also set the driver, Trusted_Connection and the SPNs (all keys are optional), for example:
       {"DefaultServer": "SERVER_A", "DefaultDatabase": "MyDatabase",
        "DefaultDriver": "ODBC Driver 17 for SQL Server", "Trusted_Connection": "yes",
        "SPNs": {"SERVER_A": "MSSQLSvc/server-a.example.com:instance"}}
    Without a file, the driver is ODBC Driver 17 for SQL Server and Trusted_Connection is yes.
    On macOS it authenticates with the Kerberos ticket cache (a valid ticket is required, see klist)
    and needs the SPN of the server, so a file with SPNs is required there. On other systems the SPN
    is not used. On macOS and Linux the file can also give "ODBCSYSINI" (folder of odbcinst.ini,
    where the driver is registered) and "KRB5_CONFIG" (krb5.conf file); they are set as environment
    variables before the first connection unless already set. On Windows they are ignored.

    Arg:
       Server: Server name. None uses DefaultServer of Config.
       Database: Database name. None uses DefaultDatabase of Config.
       Config: Path of the JSON configuration file; None uses no file, so Server and Database are required.

    Returns:
       SQLAlchemy engine connected to the requested database.

    Raises:
       TypeError: If Config is not None, a string or a Path.
       FileNotFoundError: If Config is given and the file does not exist, or (macOS / Linux) the
          ODBCSYSINI folder has no odbcinst.ini or the KRB5_CONFIG file does not exist.
       ValueError: If the file is invalid, Server or Database is missing (not given and no default in
          Config), or on macOS the server has no entry in SPNs.
       ImportError: If sqlalchemy or pyodbc is not installed (pip install "koco-ml[sql]").
       RuntimeError: On macOS, if klist is not available or no Kerberos ticket cache is found.
    """
    import platform
    import urllib.parse

    if Config is not None and not isinstance(Config, (str, Path)):
        raise TypeError("Config must be the path of a JSON file (string or Path) or None.")
    Settings = _LoadSqlConfig(Config) if Config is not None else {
        "DefaultServer": None, "DefaultDatabase": None, "Driver": _SQL_DEFAULT_DRIVER,
        "TrustedConnection": _SQL_DEFAULT_TRUSTED_CONNECTION, "SPNs": {}, "Environment": {},
    }

    Server = Server or Settings["DefaultServer"]
    Database = Database or Settings["DefaultDatabase"]
    Missing = [Name for Name, Value in (("Server", Server), ("Database", Database)) if not Value]
    if Missing:
        raise ValueError(
            f"{' and '.join(Missing)} must be given, or Config must be a file with "
            f"{' and '.join('Default' + Name for Name in Missing)}. Example content:\n{_SQL_CONFIG_EXAMPLE}"
        )

    IsMac = platform.system() == "Darwin"
    SPN = Settings["SPNs"].get(Server)
    if IsMac and not SPN:
        raise ValueError(
            f'On macOS (Kerberos) server {Server!r} needs an SPN: pass Config, a JSON file with an "SPNs" entry '
            f"for it. Example content:\n{_SQL_CONFIG_EXAMPLE}"
        )
    if platform.system() != "Windows":
        # unixODBC and Kerberos read these variables when they load, so set them before pyodbc.
        _SetSqlEnvironment(Settings["Environment"])

    try:
        import sqlalchemy as sql
    except ImportError as Error:
        raise ImportError('Connection needs sqlalchemy and pyodbc: pip install "koco-ml[sql]"') from Error

    Parts = [
        f"DRIVER={{{Settings['Driver']}}}", f"SERVER={Server}", f"DATABASE={Database}",
        f"Trusted_Connection={Settings['TrustedConnection']}",
    ]
    if IsMac:
        _SetKerberosCache()
        Parts += ["TrustServerCertificate=yes", f"ServerSPN={SPN}", "MARS_Connection=yes"]
    ConnStr = urllib.parse.quote_plus(";".join(Parts) + ";")

    Engine = sql.create_engine(f"mssql+pyodbc:///?odbc_connect={ConnStr}", fast_executemany=True)

    with Engine.connect():
        _Done(f"Connected to {Server} successfully")

    return Engine

def RunFuncInSubSample(
    Table: pd.DataFrame,
    GroupColumn: str,
    Func: Callable[..., pd.DataFrame],
    **kwargs,
) -> pd.DataFrame:
    """
    Apply a function to each subgroup and combine the results.

    Arg:
       Table: Source table to split into groups.
       GroupColumn: Column used to create subgroups.
       Func: Function called with each subgroup as its first argument; must return a DataFrame.
       kwargs: Extra keyword arguments passed to Func.

    Returns:
       Combined DataFrame with GroupColumn as the first column and one result block per subgroup.

    Raises:
       ValueError: If GroupColumn does not exist in Table or the result of Func already contains GroupColumn.
       TypeError: If Func does not return a DataFrame.
    """
    _CheckColumns(Table, [GroupColumn])

    Results = []
    for GroupValue, SubGroup in Table.groupby(GroupColumn, dropna=False):
        _Step(f"{GroupColumn}: {GroupValue} - Sample Count: {len(SubGroup)} rows...")
        Result = Func(SubGroup.reset_index(drop=True), **kwargs)
        if not isinstance(Result, pd.DataFrame):
            raise TypeError(
                f"Func must return a DataFrame — got {type(Result).__name__} for group '{GroupValue}'."
            )
        Result.insert(0, GroupColumn, GroupValue)
        Results.append(Result)

    if not Results:
        return pd.DataFrame()

    return pd.concat(Results, ignore_index=True)

def _ValidateExcelSheets(
    Sheets: dict[str, pd.DataFrame | str | list[pd.DataFrame | str]],
) -> dict[str, list[pd.DataFrame | str]]:
    """
    Validate the sheet mapping of ExportExcel before any file is written.

    Arg:
       Sheets: Sheet mapping given to ExportExcel.

    Returns:
       Mapping of sheet name to a list of DataFrames and image paths.

    Raises:
       ValueError: If Sheets is empty, a sheet name is invalid or duplicated, or a sheet or table
          exceeds the Excel row / column limits.
       TypeError: If Sheets is not a dict, or a sheet item is not a DataFrame or an image path.
    """
    if not isinstance(Sheets, dict):
        raise TypeError("Sheets must be a dict of sheet name -> table(s). The argument order is ExportExcel(Sheets, FilePath).")
    if not Sheets:
        raise ValueError("Sheets must contain at least one sheet.")

    SheetNames = [str(Name) for Name in Sheets]
    InvalidNames = [Name for Name in SheetNames if not 1 <= len(Name) <= 31 or _EXCEL_INVALID_SHEET_CHARS & set(Name)]
    if InvalidNames:
        raise ValueError(f"Invalid sheet names (1-31 characters, none of []:*?/\\): {InvalidNames}")
    if len({Name.lower() for Name in SheetNames}) != len(SheetNames):
        raise ValueError("Sheet names must be unique (Excel ignores case).")

    Validated = {}
    for SheetName, Tables in Sheets.items():
        Items = [Tables] if isinstance(Tables, (pd.DataFrame, str)) else list(Tables)
        UsedRows = 0
        for Item in Items:
            if isinstance(Item, pd.DataFrame):
                if len(Item.columns) > _EXCEL_MAX_COLUMNS:
                    raise ValueError(f"Sheet '{SheetName}': a table has more than {_EXCEL_MAX_COLUMNS:,} columns.")
                UsedRows += len(Item) + 1
                if UsedRows > _EXCEL_MAX_ROWS:
                    raise ValueError(f"Sheet '{SheetName}' exceeds the Excel limit of {_EXCEL_MAX_ROWS:,} rows.")
                UsedRows += 1
            elif not isinstance(Item, str):
                raise TypeError(f"Sheet '{SheetName}': items must be DataFrames or image paths, got {type(Item).__name__}.")
        Validated[str(SheetName)] = Items
    return Validated

def ExportExcel(
    Sheets: dict[str, pd.DataFrame | str | list[pd.DataFrame | str]],
    FilePath: str | Path | None = None,
) -> Path:
    """
    Export formatted tables and images into an Excel workbook.

    Arg:
       Sheets: Sheet mapping where each value is a DataFrame, an image path, or a list of them,
          written top to bottom.
       FilePath: Output Excel file. None writes <yymmdd>_ExportExcel.xlsx to the current folder;
          a path with an extension is the file itself; a folder (existing, or a path without an
          extension) receives the default file name. Missing folders are created.

    Returns:
       Path of the written workbook.

    Raises:
       ValueError: If Sheets is empty, a sheet name is invalid or duplicated, or a sheet exceeds
          the Excel limits.
       TypeError: If Sheets is not a dict, a sheet item is neither a DataFrame nor an image path,
          or FilePath has the wrong type.
       OSError: If writing the workbook or reading an existing image fails.

    Notes:
       A file with the same name is updated and a message is printed.
       Excel allows at most 1,048,576 rows and 16,384 columns per sheet; sheet names have 1-31
          characters and cannot contain []:*?/\\. All sheets are validated before the file is written.
       The row limit counts a header row per table and one blank row between tables; image rows are not counted.
       Tables have styled headers and alternating row colors; their index is not exported.
       Only cells within each table are formatted. Each worksheet column keeps the largest
          content-based width across its tables, bounded between 12 and 75.
       Widths are calculated by column position, so duplicate column names are supported.
       A sheet containing just one DataFrame has its header frozen.
       Missing images are reported and skipped.
    """
    Sheets = _ValidateExcelSheets(Sheets)
    OutputPath = _OutputPath(FilePath, "ExportExcel", ".xlsx")

    HeaderFill = PatternFill(fgColor="5B84B4", fill_type="solid")
    OddRowFill = PatternFill(fgColor="DCE6F1", fill_type="solid")
    EvenRowFill = PatternFill(fgColor="BCCDE0", fill_type="solid")
    HeaderFont = Font(color="FFFFFF", bold=True)
    CellAlignment = Alignment(wrap_text=True, vertical="center")
    WhiteSide = Side(color="FFFFFF", style="thin")
    WhiteBorder = Border(top=WhiteSide, bottom=WhiteSide, left=WhiteSide, right=WhiteSide)

    with pd.ExcelWriter(OutputPath, engine="openpyxl") as Writer:
        for SheetName, Tables in Sheets.items():
            _Step(f"{SheetName} sheet is preparing")

            Worksheet = Writer.book.create_sheet(SheetName)
            CurrentRow = 0
            ColumnWidths = {}

            for Table in Tables:
                if isinstance(Table, pd.DataFrame):
                    Table.to_excel(Writer, sheet_name=SheetName, startrow=CurrentRow, index=False)
                    HeaderRow = CurrentRow + 1
                    LastRow = HeaderRow + len(Table)
                    ColumnCount = len(Table.columns)

                    if ColumnCount:
                        for Row in Worksheet.iter_rows(
                            min_row=HeaderRow, max_row=LastRow, min_col=1, max_col=ColumnCount,
                        ):
                            RowNumber = Row[0].row
                            RowFill = OddRowFill if (RowNumber - HeaderRow) % 2 else EvenRowFill
                            for Cell in Row:
                                Cell.fill = HeaderFill if RowNumber == HeaderRow else RowFill
                                if RowNumber == HeaderRow:
                                    Cell.font = HeaderFont
                                Cell.alignment = CellAlignment
                                Cell.border = WhiteBorder
                        Worksheet.row_dimensions[HeaderRow].height = 24

                    for ColumnNumber, ColumnName in enumerate(Table.columns, start=1):
                        Values = Table.iloc[:, ColumnNumber - 1].dropna().astype(str)
                        MaxLength = max(len(str(ColumnName)), max((len(Value) for Value in Values), default=0))
                        ColumnWidth = min(max(MaxLength + 2, 12), 75)
                        ColumnWidths[ColumnNumber] = max(ColumnWidths.get(ColumnNumber, 0), ColumnWidth)
                        Worksheet.column_dimensions[get_column_letter(ColumnNumber)].width = ColumnWidths[ColumnNumber]

                    if len(Tables) == 1:
                        Worksheet.freeze_panes = "A2"
                    CurrentRow += len(Table) + 2

                if isinstance(Table, str):
                    ImagePath = Path(Table)
                    if ImagePath.exists():
                        ExcelImg = ExcelImage(str(ImagePath))
                        Worksheet.add_image(ExcelImg, f"A{CurrentRow+1}")
                        CurrentRow += ((ExcelImg.height // 20) + 1)  # ~20px per Excel row
                    else:
                        _Alert(f"Image not found: {ImagePath}")

    _Done(f"Saved: {OutputPath}")
    return OutputPath

def ImportExcel(
    FilePath: str | Path,
    SheetName: str | int | None = 0,
) -> pd.DataFrame | dict[str, pd.DataFrame]:
    """
    Import one sheet, or every sheet, of an Excel workbook as DataFrames.

    The first row of a sheet is the header, as written by ExportExcel. A sheet that holds several
    stacked tables is returned as one table.

    Arg:
       FilePath: Input Excel file (.xlsx or .xlsm).
       SheetName: Sheet name or zero-based position of the sheet to read; the first sheet by default.
          None reads every sheet.

    Returns:
       DataFrame of the requested sheet, or a dict of sheet name -> DataFrame when SheetName is None.

    Raises:
       FileNotFoundError: If the file does not exist.
       ValueError: If the sheet does not exist.
       OSError: If the file cannot be read.
    """
    InputPath = Path(FilePath)
    if not InputPath.is_file():
        raise FileNotFoundError(f"Excel file not found: {InputPath}")

    return pd.read_excel(InputPath, sheet_name=SheetName, engine="openpyxl")

def ExportParquet(
    Table: pd.DataFrame,
    FilePath: str | Path | None = None,
) -> Path:
    """
    Export a DataFrame to a Parquet file.

    Arg:
       Table: DataFrame to export.
       FilePath: Output file. None writes <yymmdd>_ExportParquet.parquet to the current folder; a
          path with an extension is the file itself; a folder (existing, or a path without an
          extension) receives the default file name. Missing folders are created.

    Returns:
       Path of the written file; the DataFrame index is not exported.

    Raises:
       ImportError: If neither pyarrow nor fastparquet is installed.
       TypeError: If FilePath has the wrong type.
       OSError: If the file cannot be written.

    Notes:
       A file with the same name is updated and a message is printed.
    """
    OutputPath = _OutputPath(FilePath, "ExportParquet", ".parquet")
    Table.to_parquet(OutputPath, index=False)
    _Done(f"Saved: {OutputPath}")
    return OutputPath

def ImportParquet(
    FilePath: str | Path,
) -> pd.DataFrame:
    """
    Import a DataFrame from a Parquet file.

    Arg:
       FilePath: Input Parquet file path.

    Returns:
       DataFrame read from the file.

    Raises:
       ImportError: If neither pyarrow nor fastparquet is installed.
       OSError: If the file cannot be read.
    """
    return pd.read_parquet(FilePath)

def ExportJoblib(
    Data: dict,
    FilePath: str | Path | None = None,
) -> Path:
    """
    Export a dictionary to a joblib file.

    Arg:
       Data: Dictionary to export.
       FilePath: Output file. None writes <yymmdd>_ExportJoblib.joblib to the current folder; a
          path with an extension is the file itself; a folder (existing, or a path without an
          extension) receives the default file name. Missing folders are created.

    Returns:
       Path of the written file.

    Raises:
       TypeError: If FilePath has the wrong type.
       OSError: If the file cannot be written.

    Notes:
       A file with the same name is updated and a message is printed.
    """
    import joblib

    OutputPath = _OutputPath(FilePath, "ExportJoblib", ".joblib")
    joblib.dump(Data, OutputPath)
    _Done(f"Saved: {OutputPath}")
    return OutputPath

def ImportJoblib(
    FilePath: str | Path,
) -> dict:
    """
    Import a joblib file.

    Arg:
       FilePath: Input joblib file path.

    Returns:
       Dictionary read from the file.

    Raises:
       OSError: If the file cannot be read.

    Notes:
       Only load trusted files. Loading joblib files can execute arbitrary code.
    """
    import joblib

    return joblib.load(FilePath)

def ExportJson(
    Data: dict,
    FilePath: str | Path | None = None,
) -> Path:
    """
    Export a dictionary, such as a fitted plan, to a JSON file.

    Arg:
       Data: JSON-serializable dictionary to export.
       FilePath: Output file. None writes <yymmdd>_ExportJson.json to the current folder; a path
          with an extension is the file itself; a folder (existing, or a path without an
          extension) receives the default file name. Missing folders are created.

    Returns:
       Path of the written file.

    Raises:
       TypeError: If Data is not a dict or is not JSON-serializable, or FilePath has the wrong type.
       ValueError: If Data contains NaN or inf values, which are not valid JSON.
       OSError: If the file cannot be written.

    Notes:
       A file with the same name is updated and a message is printed.
    """
    if not isinstance(Data, dict):
        raise TypeError("Data must be a dict.")

    # Serialize first so a failure does not leave a truncated file behind.
    Content = json.dumps(Data, indent=2, ensure_ascii=False, allow_nan=False)
    OutputPath = _OutputPath(FilePath, "ExportJson", ".json")
    OutputPath.write_text(Content, encoding="utf-8")

    _Done(f"Saved: {OutputPath}")
    return OutputPath

def ImportJson(
    FilePath: str | Path,
) -> dict:
    """
    Import a dictionary, such as a fitted plan, from a JSON file.

    Arg:
       FilePath: Path of a file written by ExportJson.

    Returns:
       Dictionary read from the file; the Apply functions validate plans themselves.

    Raises:
       ValueError: If the top level of the file is not a JSON object.
       OSError: If the file cannot be read.
       json.JSONDecodeError: If the file is not valid JSON.
    """
    with open(FilePath, encoding="utf-8") as File:
        Data = json.load(File)

    if not isinstance(Data, dict):
        raise ValueError("The JSON file must contain an object at the top level.")

    return Data

def _PackageVersions(
    Package: str,
) -> dict:
    """
    Collect the versions recorded next to saved models.

    Arg:
       Package: Top-level package of the fitted model, e.g. "sklearn".

    Returns:
       Dict of package name to version for python, pandas, numpy and Package.

    Raises:
       ImportError: If Package is not installed.
    """
    Versions = {"python": sys.version.split()[0], "pandas": pd.__version__, "numpy": np.__version__}
    Versions[Package] = getattr(importlib.import_module(Package), "__version__", None)
    return Versions

def ExportModelArtifacts(
    ModelArtifacts: dict,
    Explanation: str = "Model Artifacts",
    DirPath: str | Path | None = None,
) -> Path:
    """
    Save model artifacts and a JSON manifest to a directory.

    Arg:
       ModelArtifacts: Mapping of algorithm names to fitted artifact dictionaries.
       Explanation: Short description recorded in the manifest.
       DirPath: Destination directory. None uses the current folder; a missing folder is created.

    Returns:
       Path of the directory. It holds one <Algorithm>.joblib file per model and Manifest.json.

    Raises:
       TypeError: If DirPath has the wrong type.
       OSError: If an artifact or the manifest cannot be written.
       KeyError: If an artifact does not contain a fitted Model.

    Notes:
       Files with the same name are updated and a message is printed for each.
    """
    import joblib

    if DirPath is not None and not isinstance(DirPath, (str, Path)):
        raise TypeError("DirPath must be a string, a Path or None.")
    DirPath = Path(DirPath) if DirPath is not None else Path()
    DirPath.mkdir(parents=True, exist_ok=True)

    Manifest = {
        "Explanation": Explanation,
        "SavedAt": datetime.now().isoformat(timespec="seconds"),
        "Environment": {"python": sys.version.split()[0], "pandas": pd.__version__, "numpy": np.__version__},
        "Models": {},
    }

    for AlgorithmName, Artifact in ModelArtifacts.items():
        joblib.dump(Artifact, _OutputPath(DirPath / f"{AlgorithmName}.joblib", "", ".joblib"))

        Package = type(Artifact["Model"]).__module__.split(".")[0]
        Manifest["Environment"].update(_PackageVersions(Package))
        Manifest["Models"][AlgorithmName] = {
            Name: Value for Name, Value in Artifact.items() if Name != "Model"
        } | {"File": f"{AlgorithmName}.joblib", "Package": Package}

    ManifestPath = _OutputPath(DirPath / "Manifest.json", "", ".json")
    ManifestPath.write_text(json.dumps(Manifest, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    _Done(f"{len(ModelArtifacts)} model artifacts saved → {DirPath}")
    return DirPath

def _FileSha256(
    FilePath: Path,
) -> str:
    """
    Calculate the SHA-256 digest of a file.

    Arg:
       FilePath: File to read.

    Returns:
       Hexadecimal digest.

    Raises:
       OSError: If the file cannot be read.
    """
    Digest = hashlib.sha256()
    with open(FilePath, "rb") as File:
        for Chunk in iter(lambda: File.read(1 << 20), b""):
            Digest.update(Chunk)
    return Digest.hexdigest()

def ExportModelBundle(
    Artifact: dict,
    Plans: dict[str, dict] | None = None,
    FilePath: str | Path | None = None,
    Explanation: str = "Model Bundle",
) -> tuple[Path, str]:
    """
    Save one model together with its preprocessing plans in a single joblib file for production.

    Arg:
       Artifact: One entry of the ModelArtifacts dict returned by RegressionModelling or
          ClassificationModelling, e.g. ModelArtifacts["RandFR"].
       Plans: Fitted plans keyed by "Null" (FitNullPlan), "Outlier" (outlier plan), "WOE" (FitWOEPlan)
          and "Scaling" (FitScalingPlan). All are optional; leave out the plans the model does not need.
       FilePath: Output file. None writes <yymmdd>_ExportModelBundle.joblib to the current folder; a
          path with an extension is the file itself; a folder (existing, or a path without an
          extension) receives the default file name. Missing folders are created.
       Explanation: Short description stored in the bundle.

    Returns:
       Tuple of (FilePath, Sha256). Sha256 is the digest of the written file; keep it to verify the
       file later with ImportModelBundle(ExpectedHash=...).

    Raises:
       ValueError: If Artifact is not a model artifact, or a plan is invalid.
       TypeError: If Plans is not a dict or FilePath has the wrong type.
       OSError: If the file cannot be written.

    Notes:
       The model is stored as a block of bytes with its own SHA-256 digest, which ImportModelBundle
          checks. The environment versions and the koco version are recorded as well.
       A file with the same name is updated and a message is printed.
       Only load trusted bundles. Loading joblib files can execute arbitrary code.
    """
    import joblib

    from . import __version__

    _CheckArtifact(Artifact)
    if Plans is None:
        Plans = {}
    if not isinstance(Plans, dict):
        raise TypeError("Plans must be a dict such as {'Null': NullPlan, 'Scaling': ScalePlan}.")
    UnknownPlans = [Name for Name in Plans if Name not in BUNDLE_PLANS]
    if UnknownPlans:
        raise ValueError(f"Unknown plan names: {UnknownPlans}. Must be some of {list(BUNDLE_PLANS)}.")
    Plans = {Name: Plans[Name] for Name in BUNDLE_PLANS if Plans.get(Name) is not None}
    for Plan in Plans.values():
        _ValidatePlan(Plan)

    Buffer = io.BytesIO()
    joblib.dump(Artifact["Model"], Buffer)
    ModelBytes = Buffer.getvalue()
    Package = type(Artifact["Model"]).__module__.split(".")[0]

    Stored = {
        "BundleFormat": 1,
        "KocoVersion": __version__,
        "CreatedAt": datetime.now().isoformat(timespec="seconds"),
        "Explanation": Explanation,
        "Environment": _PackageVersions(Package),
        "Plans": Plans,
        "Artifact": {Name: Value for Name, Value in Artifact.items() if Name != "Model"},
        "ModelBytes": ModelBytes,
        "ModelSha256": hashlib.sha256(ModelBytes).hexdigest(),
    }

    OutputPath = _OutputPath(FilePath, "ExportModelBundle", ".joblib")
    joblib.dump(Stored, OutputPath)
    FileSha256 = _FileSha256(OutputPath)

    _Done(f"Model bundle saved: {OutputPath} (plans: {list(Plans) or 'none'})")
    _Info(f"SHA-256: {FileSha256}")
    return OutputPath, FileSha256

def ImportModelBundle(
    FilePath: str | Path,
    ExpectedHash: str | None = None,
) -> dict:
    """
    Import a bundle written by ExportModelBundle and check its integrity.

    Arg:
       FilePath: Path of the bundle file.
       ExpectedHash: SHA-256 digest returned by ExportModelBundle. If given, the file is verified
          before anything is loaded.

    Returns:
       Bundle dict with the keys KocoVersion, CreatedAt, Explanation, Environment, Plans (only the
       plans that were exported), Artifact (the model artifact including its fitted Model) and
       ModelSha256. Pass it to ApplyModelBundle.

    Raises:
       ValueError: If a digest does not match, or the file is not a model bundle.
       OSError: If the file cannot be read.

    Notes:
       A message is printed when the installed package versions differ from the ones the bundle
          was saved with.
       Only load trusted bundles. Loading joblib files can execute arbitrary code.
    """
    import joblib

    from . import __version__

    if ExpectedHash is not None and _FileSha256(Path(FilePath)) != ExpectedHash.strip().lower():
        raise ValueError("The SHA-256 digest of the file does not match ExpectedHash; the file was changed or is not the exported one.")

    Stored = joblib.load(FilePath)
    if not isinstance(Stored, dict) or Stored.get("BundleFormat") != 1 or not {"Artifact", "Plans", "ModelBytes", "ModelSha256"} <= Stored.keys():
        raise ValueError("The file is not a koco model bundle.")
    if hashlib.sha256(Stored["ModelBytes"]).hexdigest() != Stored["ModelSha256"]:
        raise ValueError("The model inside the bundle does not match its SHA-256 digest; the file is corrupted.")

    Bundle = {Name: Value for Name, Value in Stored.items() if Name != "ModelBytes"}
    Bundle["Artifact"] = {"Model": joblib.load(io.BytesIO(Stored["ModelBytes"])), **Stored["Artifact"]}

    if Stored["KocoVersion"] != __version__:
        _Alert(f"koco version differs: bundle {Stored['KocoVersion']}, installed {__version__}.")
    for Package, SavedVersion in Stored["Environment"].items():
        InstalledVersion = sys.version.split()[0] if Package == "python" else getattr(importlib.import_module(Package), "__version__", None)
        if InstalledVersion != SavedVersion:
            _Alert(f"{Package} version differs: bundle {SavedVersion}, installed {InstalledVersion}.")

    return Bundle
